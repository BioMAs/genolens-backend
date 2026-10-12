"""
CosmOmics — reader of the SciLicium claim workbooks.

The "Claims_Pathway_Review_*.xlsx" workbooks bundled in app/data/claims/ are the curated
source from which `scripts/build_claim_subgroups.py` pre-fills `claim_subgroups.yaml`, the
referential the CosmOmics nova reads. This module only parses them. Nothing is written to the
database any more: the former Skin claims module imported them into `claim_pathway_mappings`,
a table the nova never reads.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Union

import openpyxl

# Default location of the bundled referential workbooks.
DEFAULT_CLAIMS_DIR = Path(__file__).resolve().parent.parent / "data" / "claims"

_PATHWAYS_SHEET = "Pathways"

# Header of the Pathways sheet -> key of the parsed row.
_PATHWAY_COLS = {
    "Term_ID": "term_id",
    "Description": "description",
    "Common_abbrev": "common_abbrev",
    "Suggested_acronym": "suggested_acronym",
    "Original_claims": "original_claims",
    "Original_direction": "original_direction",
    "Updated_claim_framing": "updated_claim_framing",
    "Updated_direction": "updated_direction",
    "Category": "category",
    "Evidence_level": "evidence_level",
    "Rationale": "rationale",
    "Caveats": "caveats",
    "RefCat": "ref_cat",
}


def _clean(value) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def parse_workbook(source: Union[Path, str]) -> list[dict]:
    """The rows of one workbook's "Pathways" sheet, keyed as in `_PATHWAY_COLS`."""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows_out: list[dict] = []
    if _PATHWAYS_SHEET in wb.sheetnames:
        ws = wb[_PATHWAYS_SHEET]
        rows = ws.iter_rows(values_only=True)
        header = [(_clean(h) or "") for h in next(rows, [])]
        idx = {name: header.index(name) for name in _PATHWAY_COLS if name in header}
        for row in rows:
            term_id = _clean(row[idx["Term_ID"]]) if "Term_ID" in idx else None
            if not term_id:
                continue
            rows_out.append(
                {attr: _clean(row[idx[col]]) for col, attr in _PATHWAY_COLS.items() if col in idx}
            )
    wb.close()
    return rows_out
