"""
Self-service pipeline inputs are normalised to TSV before R reads them.

run_multimethod_pipeline.R reads the count matrix, the sample sheet and the
comparisons file with read_tsv. The upload endpoint accepts CSV, Excel and .txt,
so a comma-separated or Excel file used to reach R as-is, parse as a single
column and fail the analysis. These tests pin the conversion helper and check
that the R command receives TSV paths.
"""

import csv
import subprocess
import sys
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pandas as pd
import pytest

from app.services.pipeline_inputs import (
    PipelineInputError,
    detect_delimiter,
    normalise_to_tsv,
    prepare_pipeline_inputs,
)

COUNTS_ROWS = [
    ["gene_id", "S1", "S2", "S3"],
    ["ENSG00000141510", "12", "0", "7"],
    ["ENSG00000012048", "300", "281", "299"],
    ["MARCH1", "5", "4", "3"],
]


def _read_tsv_rows(path) -> list[list[str]]:
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.reader(fh, delimiter="\t"))


def _write_delimited(path, rows, delimiter):
    with open(path, "w", encoding="utf-8", newline="") as fh:
        csv.writer(fh, delimiter=delimiter, lineterminator="\n").writerows(rows)
    return path


# ─────────────────────────────────────────────────────────────────────────────
# normalise_to_tsv
# ─────────────────────────────────────────────────────────────────────────────


def test_csv_is_rewritten_as_tsv_with_header_and_ids_untouched(tmp_path):
    src = _write_delimited(tmp_path / "counts.csv", COUNTS_ROWS, ",")
    dest = tmp_path / "out" / "counts.tsv"

    result = normalise_to_tsv(src, dest)

    assert result == dest
    assert _read_tsv_rows(dest) == COUNTS_ROWS


def test_csv_keeps_blank_first_header_and_quoted_commas(tmp_path):
    # R's write.csv leaves the row-name header blank; free text may hold commas.
    rows = [["", "condition", "note"], ["S1", "ctrl", "batch A, day 1"]]
    src = _write_delimited(tmp_path / "samples.csv", rows, ",")
    dest = tmp_path / "samples.tsv"

    normalise_to_tsv(src, dest)

    assert _read_tsv_rows(dest) == rows
    # pandas' default reader would have renamed it "Unnamed: 0"
    assert dest.read_text(encoding="utf-8").startswith("\tcondition\tnote\n")


def test_xlsx_first_sheet_is_rewritten_as_tsv(tmp_path):
    src = tmp_path / "counts.xlsx"
    counts = pd.DataFrame({"gene_id": ["ENSG00000141510", "MARCH1"], "S1": [12, 5], "S2": [0, 4]})
    with pd.ExcelWriter(src) as writer:
        counts.to_excel(writer, sheet_name="counts", index=False)
        pd.DataFrame({"ignored": [1]}).to_excel(writer, sheet_name="notes", index=False)
    dest = tmp_path / "counts.tsv"

    normalise_to_tsv(src, dest)

    assert _read_tsv_rows(dest) == [
        ["gene_id", "S1", "S2"],
        ["ENSG00000141510", "12", "0"],  # integers stay integers, not "12.0"
        ["MARCH1", "5", "4"],
    ]


def test_xlsx_drops_empty_rows_and_trailing_empty_columns(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["sample_id", "condition"])
    ws.append(["S1", "ctrl"])
    ws.append([])
    ws.append(["S2", "treated"])
    ws["D5"] = ""  # stray formatted cell far to the right
    ws["D5"].number_format = "0.00"
    src = tmp_path / "samples.xlsx"
    wb.save(src)
    dest = tmp_path / "samples.tsv"

    normalise_to_tsv(src, dest)

    assert _read_tsv_rows(dest) == [
        ["sample_id", "condition"],
        ["S1", "ctrl"],
        ["S2", "treated"],
    ]


def test_tsv_is_passed_through_untouched(tmp_path):
    src = _write_delimited(tmp_path / "counts.tsv", COUNTS_ROWS, "\t")
    before = src.read_bytes()
    dest = tmp_path / "out" / "counts.tsv"

    result = normalise_to_tsv(src, dest)

    assert result == src
    assert not dest.exists()
    assert src.read_bytes() == before


def test_tsv_with_commas_in_cells_stays_tsv(tmp_path):
    rows = [["sample_id", "condition", "description"], ["S1", "ctrl", "liver, 24h, male"]]
    src = _write_delimited(tmp_path / "samples.tsv", rows, "\t")

    assert normalise_to_tsv(src, tmp_path / "x.tsv") == src


def test_comma_separated_txt_is_rewritten_as_tsv(tmp_path):
    src = _write_delimited(tmp_path / "counts.txt", COUNTS_ROWS, ",")
    dest = tmp_path / "counts.tsv"

    assert normalise_to_tsv(src, dest) == dest
    assert _read_tsv_rows(dest) == COUNTS_ROWS


def test_tab_separated_txt_is_passed_through(tmp_path):
    src = _write_delimited(tmp_path / "counts.txt", COUNTS_ROWS, "\t")

    assert normalise_to_tsv(src, tmp_path / "counts.tsv") == src


def test_semicolon_csv_from_european_excel_is_rewritten(tmp_path):
    rows = [["name", "numerator", "denominator"], ["traité_vs_ctrl", "traité", "ctrl"]]
    src = tmp_path / "comparisons.csv"
    src.write_bytes(("\r\n".join(";".join(r) for r in rows) + "\r\n").encode("cp1252"))
    dest = tmp_path / "comparisons.tsv"

    normalise_to_tsv(src, dest)

    assert _read_tsv_rows(dest) == rows  # cp1252 accents re-encoded as UTF-8


def test_utf8_bom_is_not_kept_in_the_first_header(tmp_path):
    src = tmp_path / "counts.csv"
    src.write_bytes("﻿gene_id,S1\nTP53,3\n".encode("utf-8"))
    dest = tmp_path / "counts.tsv"

    normalise_to_tsv(src, dest)

    assert _read_tsv_rows(dest)[0] == ["gene_id", "S1"]


def test_converted_output_reads_back_as_the_same_table(tmp_path):
    src = _write_delimited(tmp_path / "counts.csv", COUNTS_ROWS, ",")
    dest = tmp_path / "counts.tsv"
    normalise_to_tsv(src, dest)

    df = pd.read_csv(dest, sep="\t", dtype=str)

    assert list(df.columns) == COUNTS_ROWS[0]
    assert df["gene_id"].tolist() == [r[0] for r in COUNTS_ROWS[1:]]


def test_upload_whitelist_matches_the_wizard_formats():
    # The setup wizard offers .xls; the upload endpoint used to reject it.
    from app.core.config import Settings

    default = Settings.model_fields["ALLOWED_FILE_EXTENSIONS"].default
    assert set(default) == {".csv", ".tsv", ".txt", ".xlsx", ".xls"}


# ─────────────────────────────────────────────────────────────────────────────
# detect_delimiter
# ─────────────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "text, expected",
    [
        ("a\tb\n1\t2\n", "\t"),
        ("a,b\n1,2\n", ","),
        ("a;b\n1,5;2,3\n", ";"),  # decimal commas do not fool it
        ("s1\ts2\ng\t1\t2\n", "\t"),  # write.table header one field short
    ],
)
def test_detect_delimiter(text, expected):
    assert detect_delimiter(text) == expected


def test_detect_delimiter_rejects_empty_text():
    with pytest.raises(PipelineInputError):
        detect_delimiter("\n\n")


# ─────────────────────────────────────────────────────────────────────────────
# prepare_pipeline_inputs
# ─────────────────────────────────────────────────────────────────────────────


def test_prepare_pipeline_inputs_returns_tsv_for_every_input(tmp_path):
    counts = _write_delimited(tmp_path / "counts.csv", COUNTS_ROWS, ",")
    samples = _write_delimited(
        tmp_path / "samples.tsv", [["sample_id", "condition"], ["S1", "ctrl"]], "\t"
    )
    comparisons = _write_delimited(
        tmp_path / "comparisons.txt",
        [["name", "numerator", "denominator"], ["a_vs_b", "a", "b"]],
        ",",
    )
    workdir = tmp_path / "inputs"

    out = prepare_pipeline_inputs(
        {"counts": str(counts), "samples": str(samples), "comparisons": str(comparisons)},
        workdir,
    )

    assert out == {
        "counts": str(workdir / "counts.tsv"),
        "samples": str(samples),
        "comparisons": str(workdir / "comparisons.tsv"),
    }


def test_prepare_pipeline_inputs_names_the_file_that_failed(tmp_path):
    empty = tmp_path / "samples.csv"
    empty.write_text("\n")

    with pytest.raises(PipelineInputError, match=r"sample sheet \(samples\.csv\)"):
        prepare_pipeline_inputs({"samples": str(empty)}, tmp_path / "inputs")


def test_prepare_pipeline_inputs_reports_missing_file(tmp_path):
    with pytest.raises(PipelineInputError, match="count matrix"):
        prepare_pipeline_inputs({"counts": str(tmp_path / "nope.csv")}, tmp_path / "inputs")


# ─────────────────────────────────────────────────────────────────────────────
# run_self_service_analysis → the R command receives TSV
# ─────────────────────────────────────────────────────────────────────────────


def _legacy_tasks_module():
    # app/worker/tasks/ shadows tasks.py and loads it under this name.
    import app.worker.tasks  # noqa: F401

    return sys.modules["app.worker._tasks_legacy"]


def _fake_session_factory(scalar_results):
    db = MagicMock()
    db.scalar = AsyncMock(side_effect=scalar_results)
    db.commit = AsyncMock()
    db.refresh = AsyncMock()
    db.execute = AsyncMock()

    class _Ctx:
        async def __aenter__(self):
            return db

        async def __aexit__(self, *exc):
            return False

    return lambda: (lambda: _Ctx())


def test_r_command_receives_tsv_paths_for_csv_and_xlsx_uploads(tmp_path, monkeypatch):
    tasks = _legacy_tasks_module()

    storage = tmp_path / "storage"
    raw = storage / "projects" / "p1" / "raw"
    raw.mkdir(parents=True)
    _write_delimited(raw / "counts.csv", COUNTS_ROWS, ",")
    pd.DataFrame({"sample_id": ["S1", "S2", "S3"], "condition": ["a", "a", "b"]}).to_excel(
        raw / "samples.xlsx", index=False
    )
    _write_delimited(
        raw / "comparisons.tsv",
        [["name", "numerator", "denominator"], ["b_vs_a", "b", "a"]],
        "\t",
    )

    r_dir = tmp_path / "r_scripts"
    r_dir.mkdir()
    (r_dir / "run_multimethod_pipeline.R").write_text("# stub\n")

    analysis = SimpleNamespace(
        status=None,
        current_step=None,
        progress_log=[],
        params={},
        matrix_dataset_id=uuid4(),
        samples_dataset_id=uuid4(),
        comparisons_dataset_id=uuid4(),
    )
    # db.scalar is called for the analysis, then matrix, samples, comparisons
    datasets = [
        SimpleNamespace(raw_file_path="projects/p1/raw/counts.csv"),
        SimpleNamespace(raw_file_path="projects/p1/raw/samples.xlsx"),
        SimpleNamespace(raw_file_path="projects/p1/raw/comparisons.tsv"),
    ]

    seen = {}

    def fake_run(cmd, **kwargs):
        # Read the inputs now: the task deletes its working dir afterwards.
        args = dict(zip(cmd[2::2], cmd[3::2]))
        for flag in ("--counts", "--samples", "--comparisons"):
            with open(args[flag], encoding="utf-8") as fh:
                seen[flag] = (args[flag], fh.readline())
        return subprocess.CompletedProcess(cmd, returncode=1, stdout="", stderr="stub R")

    monkeypatch.setattr(tasks.settings, "LOCAL_STORAGE_PATH", str(storage))
    monkeypatch.setattr(tasks, "R_SCRIPTS_DIR", r_dir)
    monkeypatch.setattr(tasks, "_make_worker_session", _fake_session_factory([analysis, *datasets]))
    monkeypatch.setattr(subprocess, "run", fake_run)

    analysis_id = str(uuid4())
    result = tasks.run_self_service_analysis(analysis_id)

    # The stub R exits 1, so the task fails — but only after R was called.
    assert result["status"] == "failed"
    assert "R pipeline exited with code 1" in result["error"]

    inputs_dir = f"/tmp/analyses/{analysis_id}/inputs"
    assert seen["--counts"] == (f"{inputs_dir}/counts.tsv", "gene_id\tS1\tS2\tS3\n")
    assert seen["--samples"] == (f"{inputs_dir}/samples.tsv", "sample_id\tcondition\n")
    # Already TSV: handed over from storage untouched
    assert seen["--comparisons"] == (
        str(raw / "comparisons.tsv"),
        "name\tnumerator\tdenominator\n",
    )
    assert any(
        entry["step"] == "preparing_inputs"
        and entry["message"] == "Converted to TSV: count matrix, sample sheet"
        for entry in analysis.progress_log
    )
